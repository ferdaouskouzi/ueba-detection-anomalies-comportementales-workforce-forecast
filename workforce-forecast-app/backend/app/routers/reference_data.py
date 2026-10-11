"""
Référentiels métier réels — collaborateurs, projets, codes d'activité, et les imports
.xlsx qui les alimentent (cahier des charges §2.3, §3.2, §3.3). Ne touche pas à
l'authentification (voir routers/auth.py).
"""
import io
import json
import re

import openpyxl
from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, Request, UploadFile, File, status
from sqlalchemy.orm import Session

from ..audit import log_read_access
from ueba.tasks import planifier_rescore
from ..database import get_db
from ..deps import get_current_user, require_roles
from ..models import AppUser, Employee, Project, ActivityCode, Direction, Department, ImportBatch, Role
from ..readiness import department_is_ready
from ..schemas import (
    EmployeeOut, ProjectOut, ActivityCodeOut, ImportReport, DepartmentImportSummary,
    ImportHistoryOut, DirectionOut, DepartmentOut,
)
from ..import_logic import validate_employees, validate_projects, process_mapping, resolve_hierarchy

router = APIRouter(tags=["reference-data"])


def _department_summaries(db: Session) -> list[DepartmentImportSummary]:
    """Photo de tous les départements — utilisée pour que l'admin sache, après un import,
    lesquels sont prêts pour leurs leads (cahier des charges §3.3)."""
    summaries = []
    for dept in db.query(Department).order_by(Department.nom).all():
        count = db.query(Employee).filter(Employee.department_id == dept.id).count()
        summaries.append(DepartmentImportSummary(
            department_id=dept.id, department_nom=dept.nom,
            collaborateurs=count, pret=department_is_ready(db, dept.id),
        ))
    return summaries


def _read_sheet(wb, sheet_name: str) -> list[dict]:
    if sheet_name not in wb.sheetnames:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Feuille '{sheet_name}' absente du fichier")
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(h).strip().lower() if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        return []
    rows = []
    for raw in rows_iter:
        if all(v is None for v in raw):
            continue  # ligne vide en fin de feuille
        rows.append({header[i]: raw[i] for i in range(min(len(header), len(raw)))})
    return rows


def _slugify_id(name: str, prefix: str) -> str:
    slug = re.sub(r"[^A-Z0-9]+", "_", name.strip().upper()).strip("_")
    return f"{prefix}_{slug}" if slug else f"{prefix}_INCONNU"


def _get_or_create_direction(db: Session, nom: str) -> Direction:
    d = db.query(Direction).filter(Direction.nom == nom).first()
    if d is None:
        d = Direction(id=_slugify_id(nom, "DIR"), nom=nom)
        db.add(d)
        db.flush()
    return d


def _get_or_create_department(db: Session, nom: str, direction_id: str) -> Department:
    d = db.query(Department).filter(Department.nom == nom).first()
    if d is None:
        d = Department(id=_slugify_id(nom, "DEP"), nom=nom, direction_id=direction_id,
                        lead_pernr="", lead_titre="")
        db.add(d)
        db.flush()
    return d


@router.post("/admin/import/employees", response_model=ImportReport)
def import_employees(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: AppUser = Depends(require_roles("ADMIN")),
):
    content = file.file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Fichier .xlsx illisible")

    people_rows = _read_sheet(wb, "people")
    mapping_rows = _read_sheet(wb, "mapping")

    known_pernrs = {p for (p,) in db.query(Employee.pernr).all()}
    people_report = validate_employees(people_rows, known_pernrs)

    for r in people_report:
        if r["statut"] != "accepté":
            continue
        emp = db.query(Employee).filter(Employee.pernr == r["pernr"]).first()
        if emp is None:
            emp = Employee(pernr=r["pernr"])
            db.add(emp)
        emp.nom = r["nom"]
        emp.manager_pernr = r["manager_pernr"]
        emp.statut = r["statut_employe"]
        emp.grade = r["grade"]
    db.flush()

    # Jointure mapping.name -> pernr par nom exact (fragile en cas d'homonymes, cf. doc §2.3)
    people_by_name: dict[str, list[str]] = {}
    for pernr, nom in db.query(Employee.pernr, Employee.nom).all():
        people_by_name.setdefault(nom, []).append(pernr)

    mapping_report = process_mapping(mapping_rows, people_by_name)

    for r in mapping_report:
        if r["statut"] != "accepté":
            continue
        emp = db.query(Employee).filter(Employee.pernr == r["pernr"]).first()
        role = r["role"]
        direction = _get_or_create_direction(db, r["direction"])
        if role == "Head":
            direction.head_pernr = emp.pernr
            emp.direction_id = direction.id
        elif role in ("EUM", "EUD"):
            dept = _get_or_create_department(db, r["department"], direction.id)
            dept.lead_pernr = emp.pernr
            dept.lead_titre = role
            emp.department_id = dept.id
            emp.direction_id = direction.id
        else:
            r["statut"] = "rejeté"
            r["motif"] = "Rôle mapping inconnu"
            r["detail"] = role
    db.flush()

    # Propagation department/direction du reste de la hiérarchie via manager_pernr
    all_employees = {
        e.pernr: {"manager_pernr": e.manager_pernr, "department_id": e.department_id, "direction_id": e.direction_id}
        for e in db.query(Employee).all()
    }
    sans_departement = resolve_hierarchy(all_employees)
    for pernr, data in all_employees.items():
        emp = db.query(Employee).filter(Employee.pernr == pernr).first()
        emp.department_id = data["department_id"]
        emp.direction_id = data["direction_id"]

    full_report = people_report + mapping_report
    acceptees = sum(1 for r in full_report if r["statut"] == "accepté")
    rejetees = len(full_report) - acceptees

    batch = ImportBatch(
        kind="employees", filename=file.filename or "employees.xlsx", imported_by_pernr=user.pernr,
        lignes_acceptees=acceptees, lignes_rejetees=rejetees,
        rapport_json=json.dumps({"rapport": full_report, "sans_departement_resolu": sans_departement}, ensure_ascii=False),
    )
    db.add(batch)
    db.commit()

    return ImportReport(
        kind="employees", total_lignes=len(people_rows) + len(mapping_rows),
        acceptees=acceptees, rejetees=rejetees, rapport=full_report,
        par_departement=_department_summaries(db),
    )


@router.post("/admin/import/projects", response_model=ImportReport)
def import_projects(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: AppUser = Depends(require_roles("ADMIN")),
):
    content = file.file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Fichier .xlsx illisible")

    sheet_name = "projects" if "projects" in wb.sheetnames else wb.sheetnames[0]
    rows = _read_sheet(wb, sheet_name)
    report = validate_projects(rows)

    for r in report:
        if r["statut"] != "accepté":
            continue
        proj = db.query(Project).filter(Project.code == r["code"]).first()
        if proj is None:
            proj = Project(code=r["code"])
            db.add(proj)
        proj.nom = r["nom"]
        proj.client = r["client"]
        proj.statut = r["statut_projet"]

    acceptees = sum(1 for r in report if r["statut"] == "accepté")
    rejetees = len(report) - acceptees

    batch = ImportBatch(
        kind="projects", filename=file.filename or "projects.xlsx", imported_by_pernr=user.pernr,
        lignes_acceptees=acceptees, lignes_rejetees=rejetees,
        rapport_json=json.dumps({"rapport": report}, ensure_ascii=False),
    )
    db.add(batch)
    db.commit()

    return ImportReport(
        kind="projects", total_lignes=len(rows), acceptees=acceptees, rejetees=rejetees, rapport=report,
        par_departement=_department_summaries(db),
    )


@router.get("/directions", response_model=list[DirectionOut])
def list_directions(db: Session = Depends(get_db), user: AppUser = Depends(get_current_user)):
    return db.query(Direction).all()


@router.get("/departments", response_model=list[DepartmentOut])
def list_departments(
    direction_id: str | None = None,
    db: Session = Depends(get_db),
    user: AppUser = Depends(get_current_user),
):
    q = db.query(Department)
    if direction_id:
        q = q.filter(Department.direction_id == direction_id)
    return q.all()


@router.get("/employees", response_model=list[EmployeeOut])
def list_employees(
    request: Request,
    background_tasks: BackgroundTasks,
    department_id: str | None = None,
    db: Session = Depends(get_db),
    user: AppUser = Depends(get_current_user),
):
    q = db.query(Employee)
    # Contrôle de périmètre rejoué côté serveur (§2.2/§4.1) — jamais fait uniquement côté client.
    if user.role == Role.EUM_EUD:
        q = q.filter(Employee.department_id == user.department_id)
    elif user.role == Role.HEAD_OF_DIRECTION:
        q = q.filter(Employee.direction_id == user.direction_id)
    elif department_id:
        q = q.filter(Employee.department_id == department_id)
    result = q.all()

    ip = request.client.host if request.client else ""
    log_read_access(
        db, user, "EMPLOYEES_READ",
        f"department_id_filtre={department_id or '-'} lignes={len(result)}", ip,
        target_pernrs=[e.pernr for e in result],
    )
    # Scoring UEBA hors du chemin de la réponse (voir backend/ueba/).
    planifier_rescore(background_tasks, user.pernr)

    return result


@router.get("/projects", response_model=list[ProjectOut])
def list_projects(
    statut: str | None = None,
    db: Session = Depends(get_db),
    user: AppUser = Depends(get_current_user),
):
    q = db.query(Project)
    if statut:
        q = q.filter(Project.statut == statut)
    return q.all()


@router.get("/activity-codes", response_model=list[ActivityCodeOut])
def list_activity_codes(db: Session = Depends(get_db), user: AppUser = Depends(get_current_user)):
    return db.query(ActivityCode).all()


@router.get("/admin/import-history", response_model=list[ImportHistoryOut])
def import_history(db: Session = Depends(get_db), user: AppUser = Depends(require_roles("ADMIN"))):
    return db.query(ImportBatch).order_by(ImportBatch.imported_at.desc()).all()
